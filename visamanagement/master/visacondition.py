from django.shortcuts import render
from  .models  import *
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from .serializers import  *
from django.core.paginator import Paginator
from django.db.models import Q
import uuid
from rest_framework.permissions import IsAuthenticated ,AllowAny ,BasePermission 
from django.shortcuts import get_object_or_404
from .pagination import  *
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework.exceptions import ValidationError
from rest_framework_simplejwt.exceptions import TokenError  
from rest_framework_simplejwt.authentication import JWTAuthentication
from import_export.formats.base_formats import CSV, XLSX
from tablib import Dataset
import openpyxl
from django.http import HttpResponse
from uuid import UUID
from datetime import datetime  
from django.db import IntegrityError,transaction
import csv
import io
import pytz
from django.utils import timezone
import unicodedata


india_tz = pytz.timezone('Asia/Kolkata')

class IsAdminUser(BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.is_superuser)
    




class WorkRightsListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = WorkRights.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = WorkRightsSerializer(result_page, many=True)

        return paginator.get_paginated_response(serializer.data)


class WorkRightsCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = WorkRights.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "WorkRight with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = WorkRightsSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "WorkRight created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)


class WorkRightsRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            workright = WorkRights.objects.get(uuid=uuid, is_deleted=False)
        except WorkRights.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "WorkRight not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = WorkRightsSerializer(workright)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "WorkRight retrieved successfully",
            "data": serializer.data
        })


class WorkRightsUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            workright = WorkRights.objects.get(uuid=uuid, is_deleted=False)
        except WorkRights.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "WorkRight not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = WorkRightsSerializer(workright, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "WorkRight updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)


class WorkRightsDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                workright = WorkRights.objects.get(uuid=uuid)
                workright.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "WorkRight permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except WorkRights.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "WorkRight not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            workrights = WorkRights.objects.all()
            count = workrights.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No WorkRights found to delete.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)
            workrights.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} WorkRights permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        if not valid_uuids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No valid UUIDs provided.",
                "data": {"invalid_uuids": invalid_uuids}
            }, status=status.HTTP_400_BAD_REQUEST)

        workrights = WorkRights.objects.filter(uuid__in=valid_uuids)
        count = workrights.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching WorkRights found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=status.HTTP_404_NOT_FOUND)

        workrights.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} WorkRight(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })


class WorkRightsExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Work Rights',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = WorkRights.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'WorkRights'

        for wr in queryset:
            row = []
            for field in field_list:
                value = getattr(wr, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'workrights.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'workrights.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


class WorkRightsImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"work rights"}  # adjust as needed
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets,
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets,
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("work rights")) if row.get("work rights") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing WorkRight name"})
                    continue

                existing = WorkRights.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "Work Rights": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    WorkRights.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)




class WorkRightsDuringStudyListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = WorkRightsDuringStudy.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = WorkRightsDuringStudySerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# -------------------------------------------------------------------
# CREATE API
# -------------------------------------------------------------------
class WorkRightsDuringStudyCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = WorkRightsDuringStudy.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "WorkRightDuringStudy with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = WorkRightsDuringStudySerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "WorkRightDuringStudy created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)


# -------------------------------------------------------------------
# RETRIEVE API
# -------------------------------------------------------------------
class WorkRightsDuringStudyRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = WorkRightsDuringStudy.objects.get(uuid=uuid, is_deleted=False)
        except WorkRightsDuringStudy.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "WorkRightDuringStudy not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = WorkRightsDuringStudySerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "WorkRightDuringStudy retrieved successfully",
            "data": serializer.data
        })


# -------------------------------------------------------------------
# UPDATE API
# -------------------------------------------------------------------
class WorkRightsDuringStudyUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = WorkRightsDuringStudy.objects.get(uuid=uuid, is_deleted=False)
        except WorkRightsDuringStudy.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "WorkRightDuringStudy not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = WorkRightsDuringStudySerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "WorkRightDuringStudy updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)


# -------------------------------------------------------------------
# DELETE API
# -------------------------------------------------------------------
class WorkRightsDuringStudyDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = WorkRightsDuringStudy.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "WorkRightDuringStudy permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except WorkRightsDuringStudy.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "WorkRightDuringStudy not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            objs = WorkRightsDuringStudy.objects.all()
            count = objs.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No WorkRightsDuringStudy found to delete.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} WorkRightsDuringStudy permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        if not valid_uuids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No valid UUIDs provided.",
                "data": {"invalid_uuids": invalid_uuids}
            }, status=status.HTTP_400_BAD_REQUEST)

        objs = WorkRightsDuringStudy.objects.filter(uuid__in=valid_uuids)
        count = objs.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching WorkRightsDuringStudy found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=status.HTTP_404_NOT_FOUND)

        objs.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} WorkRightDuringStudy(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })


# -------------------------------------------------------------------
# EXPORT API
# -------------------------------------------------------------------
class WorkRightsDuringStudyExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Work Rights During Study',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = WorkRightsDuringStudy.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'WorkRightsDuringStudy'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'workrights_during_study.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'workrights_during_study.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


# -------------------------------------------------------------------
# IMPORT API
# -------------------------------------------------------------------
class WorkRightsDuringStudyImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"work rights during study"}
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets,
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets,
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("work rights during study")) if row.get("work rights during study") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing Work Right During Study name"})
                    continue

                existing = WorkRightsDuringStudy.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "Work Rights During Study": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    WorkRightsDuringStudy.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)



class WorkRightsDuringVacationListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = WorkRightsDuringVacation.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = WorkRightsDuringVacationSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ---------------- CREATE API ----------------
class WorkRightsDuringVacationCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = WorkRightsDuringVacation.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "WorkRightDuringVacation with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = WorkRightsDuringVacationSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "WorkRightDuringVacation created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)


# ---------------- RETRIEVE API ----------------
class WorkRightsDuringVacationRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = WorkRightsDuringVacation.objects.get(uuid=uuid, is_deleted=False)
        except WorkRightsDuringVacation.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "WorkRightDuringVacation not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = WorkRightsDuringVacationSerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "WorkRightDuringVacation retrieved successfully",
            "data": serializer.data
        })


# ---------------- UPDATE API ----------------
class WorkRightsDuringVacationUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = WorkRightsDuringVacation.objects.get(uuid=uuid, is_deleted=False)
        except WorkRightsDuringVacation.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "WorkRightDuringVacation not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = WorkRightsDuringVacationSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "WorkRightDuringVacation updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)


# ---------------- DELETE API ----------------
class WorkRightsDuringVacationDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = WorkRightsDuringVacation.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "WorkRightDuringVacation permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except WorkRightsDuringVacation.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "WorkRightDuringVacation not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            objs = WorkRightsDuringVacation.objects.all()
            count = objs.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No WorkRightsDuringVacation found to delete.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} WorkRightsDuringVacation permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        if not valid_uuids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No valid UUIDs provided.",
                "data": {"invalid_uuids": invalid_uuids}
            }, status=status.HTTP_400_BAD_REQUEST)

        objs = WorkRightsDuringVacation.objects.filter(uuid__in=valid_uuids)
        count = objs.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching WorkRightsDuringVacation found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=status.HTTP_404_NOT_FOUND)

        objs.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} WorkRightDuringVacation(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })


# ---------------- EXPORT API ----------------
class WorkRightsDuringVacationExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Work Rights During Vacation',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = WorkRightsDuringVacation.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'WorkRightsDuringVacation'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'workrights_during_vacation.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'workrights_during_vacation.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


# ---------------- IMPORT API ----------------
class WorkRightsDuringVacationImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"work rights during vacation"}
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets,
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets,
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("work rights during vacation")) if row.get("work rights during vacation") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing WorkRightDuringVacation name"})
                    continue

                existing = WorkRightsDuringVacation.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "WorkRightsDuringVacation": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    WorkRightsDuringVacation.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)


class WorkRightsAfterStudyListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = WorkRightsAfterStudy.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = WorkRightsAfterStudySerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ---------------- CREATE API ----------------
class WorkRightsAfterStudyCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = WorkRightsAfterStudy.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "WorkRightAfterStudy with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = WorkRightsAfterStudySerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "WorkRightAfterStudy created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)


# ---------------- RETRIEVE API ----------------
class WorkRightsAfterStudyRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = WorkRightsAfterStudy.objects.get(uuid=uuid, is_deleted=False)
        except WorkRightsAfterStudy.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "WorkRightAfterStudy not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = WorkRightsAfterStudySerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "WorkRightAfterStudy retrieved successfully",
            "data": serializer.data
        })


# ---------------- UPDATE API ----------------
class WorkRightsAfterStudyUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = WorkRightsAfterStudy.objects.get(uuid=uuid, is_deleted=False)
        except WorkRightsAfterStudy.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "WorkRightAfterStudy not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = WorkRightsAfterStudySerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "WorkRightAfterStudy updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)


# ---------------- DELETE API ----------------
class WorkRightsAfterStudyDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = WorkRightsAfterStudy.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "WorkRightAfterStudy permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except WorkRightsAfterStudy.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "WorkRightAfterStudy not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            objs = WorkRightsAfterStudy.objects.all()
            count = objs.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No WorkRightsAfterStudy found to delete.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} WorkRightsAfterStudy permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        if not valid_uuids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No valid UUIDs provided.",
                "data": {"invalid_uuids": invalid_uuids}
            }, status=status.HTTP_400_BAD_REQUEST)

        objs = WorkRightsAfterStudy.objects.filter(uuid__in=valid_uuids)
        count = objs.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching WorkRightsAfterStudy found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=status.HTTP_404_NOT_FOUND)

        objs.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} WorkRightAfterStudy(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })


# ---------------- EXPORT API ----------------
class WorkRightsAfterStudyExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Work Rights After Study',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = WorkRightsAfterStudy.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'WorkRightsAfterStudy'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'workrights_after_study.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'workrights_after_study.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


# ---------------- IMPORT API ----------------
class WorkRightsAfterStudyImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"work rights after study"}
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets,
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets,
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("work rights after study")) if row.get("work rights after study") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing WorkRightAfterStudy name"})
                    continue

                existing = WorkRightsAfterStudy.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "WorkRightsAfterStudy": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    WorkRightsAfterStudy.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)


class PRPossibilityListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = PRPossibility.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = PRPossibilitySerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ---------------- CREATE API ----------------
class PRPossibilityCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = PRPossibility.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "PRPossibility with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = PRPossibilitySerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "PRPossibility created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)


# ---------------- RETRIEVE API ----------------
class PRPossibilityRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = PRPossibility.objects.get(uuid=uuid, is_deleted=False)
        except PRPossibility.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "PRPossibility not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = PRPossibilitySerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "PRPossibility retrieved successfully",
            "data": serializer.data
        })


# ---------------- UPDATE API ----------------
class PRPossibilityUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = PRPossibility.objects.get(uuid=uuid, is_deleted=False)
        except PRPossibility.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "PRPossibility not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = PRPossibilitySerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "PRPossibility updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)


# ---------------- DELETE API ----------------
class PRPossibilityDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = PRPossibility.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "PRPossibility permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except PRPossibility.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "PRPossibility not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            objs = PRPossibility.objects.all()
            count = objs.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No PRPossibility found to delete.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} PRPossibility permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        if not valid_uuids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No valid UUIDs provided.",
                "data": {"invalid_uuids": invalid_uuids}
            }, status=status.HTTP_400_BAD_REQUEST)

        objs = PRPossibility.objects.filter(uuid__in=valid_uuids)
        count = objs.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching PRPossibility found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=status.HTTP_404_NOT_FOUND)

        objs.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} PRPossibility(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })


# ---------------- EXPORT API ----------------
class PRPossibilityExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'PR Possibility',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = PRPossibility.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'PRPossibility'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'pr_possibility.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'pr_possibility.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


# ---------------- IMPORT API ----------------
class PRPossibilityImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"pr possibility"}
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets,
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets,
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("pr possibility")) if row.get("pr possibility") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing PR Possibility name"})
                    continue

                existing = PRPossibility.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "PR Possibility": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    PRPossibility.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)



class SpouseCanApplyListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = SpouseCanApplywithCandidate.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = SpouseCanApplywithCandidateSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ---------------- CREATE API ----------------
class SpouseCanApplyCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = SpouseCanApplywithCandidate.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "SpouseCanApplywithCandidate with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = SpouseCanApplywithCandidateSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "SpouseCanApplywithCandidate created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)


# ---------------- RETRIEVE API ----------------
class SpouseCanApplyRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = SpouseCanApplywithCandidate.objects.get(uuid=uuid, is_deleted=False)
        except SpouseCanApplywithCandidate.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "SpouseCanApplywithCandidate not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = SpouseCanApplywithCandidateSerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "SpouseCanApplywithCandidate retrieved successfully",
            "data": serializer.data
        })


# ---------------- UPDATE API ----------------
class SpouseCanApplyUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = SpouseCanApplywithCandidate.objects.get(uuid=uuid, is_deleted=False)
        except SpouseCanApplywithCandidate.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "SpouseCanApplywithCandidate not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = SpouseCanApplywithCandidateSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "SpouseCanApplywithCandidate updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)


# ---------------- DELETE API ----------------
class SpouseCanApplyDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = SpouseCanApplywithCandidate.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "SpouseCanApplywithCandidate permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except SpouseCanApplywithCandidate.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "SpouseCanApplywithCandidate not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            objs = SpouseCanApplywithCandidate.objects.all()
            count = objs.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No records found to delete.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} records permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        if not valid_uuids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No valid UUIDs provided.",
                "data": {"invalid_uuids": invalid_uuids}
            }, status=status.HTTP_400_BAD_REQUEST)

        objs = SpouseCanApplywithCandidate.objects.filter(uuid__in=valid_uuids)
        count = objs.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching records found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=status.HTTP_404_NOT_FOUND)

        objs.delete()
        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} record(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })


# ---------------- EXPORT API ----------------
class SpouseCanApplyExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Spouse Can Apply',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = SpouseCanApplywithCandidate.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'SpouseCanApplywithCandidate'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'spouse_can_apply.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'spouse_can_apply.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


# ---------------- IMPORT API ----------------
class SpouseCanApplyImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"spouse can apply"}
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets,
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets,
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("spouse can apply")) if row.get("spouse can apply") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing Name"})
                    continue

                existing = SpouseCanApplywithCandidate.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "Spouse Can Apply": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    SpouseCanApplywithCandidate.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)




class SpouseWorkRightsListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'
        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = SpouseWorkRights.objects.filter(is_deleted=False)
        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = SpouseWorkRightsSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)

class SpouseWorkRightsCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = SpouseWorkRights.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "SpouseWorkRight with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = SpouseWorkRightsSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "SpouseWorkRight created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)

class SpouseWorkRightsRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = SpouseWorkRights.objects.get(uuid=uuid, is_deleted=False)
        except SpouseWorkRights.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "SpouseWorkRight not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = SpouseWorkRightsSerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "SpouseWorkRight retrieved successfully",
            "data": serializer.data
        })


class SpouseWorkRightsUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = SpouseWorkRights.objects.get(uuid=uuid, is_deleted=False)
        except SpouseWorkRights.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "SpouseWorkRight not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = SpouseWorkRightsSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "SpouseWorkRight updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)


class SpouseWorkRightsDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = SpouseWorkRights.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "SpouseWorkRight permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except SpouseWorkRights.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "SpouseWorkRight not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        # batch delete by IDs or "all"
        if ids == "all":
            objs = SpouseWorkRights.objects.all()
            count = objs.count()
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} SpouseWorkRights permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        objs = SpouseWorkRights.objects.filter(uuid__in=valid_uuids)
        count = objs.count()
        objs.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} SpouseWorkRight(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })



class SpouseWorkRightsExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Spouse Work Rights',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        field_list = [f.strip() for f in fields.split(',')] if fields else list(field_header_map.keys())

        queryset = SpouseWorkRights.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'SpouseWorkRights'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'spouse_workrights.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'spouse_workrights.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response




class SpouseWorkRightsImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"spouse work rights"}  # adjust column name in your sheet
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets,
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets,
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("spouse work rights")) if row.get("spouse work rights") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing SpouseWorkRight name"})
                    continue

                existing = SpouseWorkRights.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "Spouse Work Rights": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    SpouseWorkRights.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)


class ChildrenCanApplyListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'
        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = ChildrenCanApplywithCandidate.objects.filter(is_deleted=False)
        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = ChildrenCanApplywithCandidateSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


class ChildrenCanApplyCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = ChildrenCanApplywithCandidate.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Record with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = ChildrenCanApplywithCandidateSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Record created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)


class ChildrenCanApplyCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = ChildrenCanApplywithCandidate.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Record with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = ChildrenCanApplywithCandidateSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Record created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)



class ChildrenCanApplyRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = ChildrenCanApplywithCandidate.objects.get(uuid=uuid, is_deleted=False)
        except ChildrenCanApplywithCandidate.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "Record not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = ChildrenCanApplywithCandidateSerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "Record retrieved successfully",
            "data": serializer.data
        })

class ChildrenCanApplyUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = ChildrenCanApplywithCandidate.objects.get(uuid=uuid, is_deleted=False)
        except ChildrenCanApplywithCandidate.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "Record not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = ChildrenCanApplywithCandidateSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Record updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        
class ChildrenCanApplyDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = ChildrenCanApplywithCandidate.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "Record permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except ChildrenCanApplywithCandidate.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "Record not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            objs = ChildrenCanApplywithCandidate.objects.all()
            count = objs.count()
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} records permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Provide a list of UUIDs in 'id' or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        objs = ChildrenCanApplywithCandidate.objects.filter(uuid__in=valid_uuids)
        count = objs.count()
        objs.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} record(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })



class ChildrenCanApplyExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Children Can Apply with Candidate',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        field_list = [f.strip() for f in fields.split(',')] if fields else list(field_header_map.keys())

        queryset = ChildrenCanApplywithCandidate.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'ChildrenCanApply'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'children_can_apply.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'children_can_apply.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response



class ChildrenCanApplyImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"children can apply with candidate"}  # column in your sheet
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames
                if not sheet_name:
                    return Response({"error": "Provide sheet_name", "available_sheets": available_sheets}, status=400)
                if sheet_name not in available_sheets:
                    return Response({"error": f'Sheet "{sheet_name}" not found', "available_sheets": available_sheets}, status=400)
                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({"statusCode": 400, "status": False, "message": f'Sheet "{sheet_name}" is empty.'}, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                if not required_headers.issubset(set(headers)):
                    return Response({"statusCode": 400, "status": False, "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"}, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({"statusCode": 400, "status": False, "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."}, status=400)
                    data.append(row_lower)
            else:
                return Response({"statusCode": 400, "status": False, "message": "Unsupported file format. Use .xlsx or .csv"}, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("children can apply with candidate")) if row.get("children can apply with candidate") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing name"})
                    continue

                existing = ChildrenCanApplywithCandidate.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "Children Can Apply with Candidate": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    ChildrenCanApplywithCandidate.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)


class ChildrenVisaCategoryListAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):

        search = request.GET.get("search", "").strip()
        visamain_uuid = request.GET.get("visamain_uuid")

        queryset = ChildrenVisaCategory.objects.filter(is_deleted=False)

        # ------------------------
        # FILTER BY VISAMAIN UUID
        # ------------------------
        if visamain_uuid:
            queryset = queryset.filter(visamain__uuid=visamain_uuid)

        # ------------------------
        # SEARCH (by visamain name or description)
        # ------------------------
        if search:
            queryset = queryset.filter(
                Q(visamain__name__icontains=search) |
                Q(description__icontains=search)
            )

        queryset = queryset.order_by("-created_at")

        serializer = ChildrenVisaCategorySerializer(queryset, many=True)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": "Children visa category list fetched successfully",
            "total": queryset.count(),
            "data": serializer.data
        })



class ChildrenVisaCategoryCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        data = request.data.copy()

        # Validate FK: VisaMain
        visamain_uuid = data.get("visamain_uuid")
        if not visamain_uuid:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "visamain_uuid is required"
            }, status=400)

        try:
            visamain_obj = VisaMain.objects.get(uuid=visamain_uuid)
        except VisaMain.DoesNotExist:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Invalid visamain_uuid"
            }, status=400)

        # Duplicate check
        desc = data.get("description", "").strip()
        existing = ChildrenVisaCategory.objects.filter(
            visamain=visamain_obj,
            description__iexact=desc,
            is_deleted=False
        ).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Category already exists"
            }, status=400)

        serializer = ChildrenVisaCategorySerializer(data=data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Children Visa Category created successfully",
                "data": serializer.data
            })

        err = " ".join([msg for msgs in serializer.errors.values() for msg in msgs])
        return Response({"statusCode": 400, "status": False, "message": err}, status=400)





class ChildrenVisaCategoryRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = ChildrenVisaCategory.objects.get(uuid=uuid, is_deleted=False)
        except ChildrenVisaCategory.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "Not found"}, status=404)

        serializer = ChildrenVisaCategorySerializer(obj)
        return Response({"statusCode": 200, "status": True, "data": serializer.data})




class ChildrenVisaCategoryUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = ChildrenVisaCategory.objects.get(uuid=uuid, is_deleted=False)
        except ChildrenVisaCategory.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "Not found"}, status=404)

        serializer = ChildrenVisaCategorySerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Updated successfully",
                "data": serializer.data
            })

        err = " ".join([msg for msgs in serializer.errors.values() for msg in msgs])
        return Response({"statusCode": 400, "status": False, "message": err}, status=400)






class ChildrenVisaCategoryDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request):
        ids = request.data.get('id')

        if not ids:
            return Response({"status": False, "message": "Provide 'id' field"}, status=400)

        if ids == "all":
            qs = ChildrenVisaCategory.objects.filter(is_deleted=False)
            count = qs.count()
            qs.update(is_deleted=True)
            return Response({"status": True, "message": f"All {count} records deleted"})

        if not isinstance(ids, list):
            return Response({"status": False, "message": "Send list of UUIDs"}, status=400)

        valid, invalid = [], []
        for u in ids:
            try:
                valid.append(UUID(u))
            except:
                invalid.append(u)

        qs = ChildrenVisaCategory.objects.filter(uuid__in=valid, is_deleted=False)
        count = qs.count()
        qs.update(is_deleted=True)

        return Response({
            "status": True,
            "message": f"{count} record(s) deleted",
            "invalid_uuids": invalid or None
        })




class ChildrenVisaCategoryExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get("format", "xlsx").lower()
        fields = request.GET.get("fields")
        uuids = request.GET.get("uuids", "").split(",")

        field_map = {
            'uuid': 'UUID',
            'visamain': 'Children Visa Main',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'created_at': 'Created At',
            'updated_at': 'Updated At',
        }

        field_list = [f.strip() for f in fields.split(',')] if fields else list(field_map.keys())

        qs = ChildrenVisaCategory.objects.filter(is_deleted=False)
        if uuids:
            qs = qs.filter(uuid__in=uuids)

        dataset = Dataset()
        dataset.headers = [field_map.get(f, f) for f in field_list]

        for obj in qs:
            row = []
            for f in field_list:
                val = getattr(obj, f, "")

                if f == "visamain" and val:
                    val = val.name
                elif f in ["created_at", "updated_at"] and val:
                    val = timezone.localtime(val).strftime("%d-%m-%Y %I:%M:%S %p")
                row.append(val)

            dataset.append(row)

        if format_type == "csv":
            response = HttpResponse(dataset.export("csv"), content_type="text/csv")
            response["Content-Disposition"] = 'attachment; filename="children_visa_category.csv"'
            return response

        file_data = io.BytesIO(dataset.export("xlsx"))
        response = HttpResponse(
            file_data.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="children_visa_category.xlsx"'
        return response





class ChildrenVisaCategoryImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=400)

        ext = file.name.split(".")[-1].lower()
        required = {"children visa main", "description"}

        parsed = []

        try:
            if ext == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                if sheet_name not in wb.sheetnames:
                    return Response({"error": "Invalid sheet name"}, status=400)
                ws = wb[sheet_name]

                headers = [str(c.value).strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
                if not required.issubset(headers):
                    return Response({"error": f"Missing headers: {required}"})

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    parsed.append((idx, dict(zip(headers, row))))

            elif ext == "csv":
                dataset = Dataset()
                dataset.load(file.read().decode("utf-8"), format="csv")

                for idx, row in enumerate(dataset.dict, start=2):
                    r = {k.strip().lower(): v for k, v in row.items()}
                    parsed.append((idx, r))

            else:
                return Response({"error": "Only CSV/XLSX supported"}, status=400)

            imported = 0
            duplicates = []
            skipped = []

            for row_num, row in parsed:
                visamain_name = str(row.get("children visa main")).strip()
                desc = str(row.get("description", "")).strip()

                if not (visamain_name and desc):
                    skipped.append({"row": row_num, "Reason": "Missing required fields"})
                    continue

                try:
                    visamain_obj = VisaMain.objects.get(name__iexact=visamain_name)
                except:
                    skipped.append({"row": row_num, "Reason": "VisaMain not found"})
                    continue

                existing = ChildrenVisaCategory.objects.filter(
                    visamain=visamain_obj,
                    description__iexact=desc
                ).first()

                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "row": row_num,
                            "Children Visa Category": visamain_name,
                            "Reason": "Duplicate entry"
                        })
                        continue
                    else:
                        existing.is_deleted = False
                        existing.save()
                        imported += 1
                else:
                    ChildrenVisaCategory.objects.create(
                        visamain=visamain_obj,
                        description=desc,
                        is_deleted=False
                    )
                    imported += 1

        except Exception as e:
            return Response({"error": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "imported_count": imported,
            "duplicates": duplicates,
            "skipped_rows": skipped,
            "message": "Import completed"
        })



class ChildrenStudyWorkRightsListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = ChildrenStudyWorkRights.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = ChildrenStudyWorkRightsSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


class ChildrenStudyWorkRightsCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = ChildrenStudyWorkRights.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "ChildrenStudyWorkRights with this name already exists."
            }, status=400)

        serializer = ChildrenStudyWorkRightsSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "ChildrenStudyWorkRights created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=400)


class ChildrenStudyWorkRightsRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = ChildrenStudyWorkRights.objects.get(uuid=uuid, is_deleted=False)
        except ChildrenStudyWorkRights.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "ChildrenStudyWorkRights not found",
                "data": None
            }, status=404)

        serializer = ChildrenStudyWorkRightsSerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "ChildrenStudyWorkRights retrieved successfully",
            "data": serializer.data
        })


class ChildrenStudyWorkRightsUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = ChildrenStudyWorkRights.objects.get(uuid=uuid, is_deleted=False)
        except ChildrenStudyWorkRights.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "ChildrenStudyWorkRights not found",
                "data": None
            }, status=404)

        serializer = ChildrenStudyWorkRightsSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "ChildrenStudyWorkRights updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=400)


class ChildrenStudyWorkRightsDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = ChildrenStudyWorkRights.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "ChildrenStudyWorkRights permanently deleted.",
                    "data": None
                }, status=204)
            except ChildrenStudyWorkRights.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "ChildrenStudyWorkRights not found.",
                    "data": None
                }, status=404)

        if ids == "all":
            objs = ChildrenStudyWorkRights.objects.all()
            count = objs.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No ChildrenStudyWorkRights found to delete.",
                    "data": None
                }, status=404)
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} ChildrenStudyWorkRights permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=400)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        if not valid_uuids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No valid UUIDs provided.",
                "data": {"invalid_uuids": invalid_uuids}
            }, status=400)

        objs = ChildrenStudyWorkRights.objects.filter(uuid__in=valid_uuids)
        count = objs.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching ChildrenStudyWorkRights found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=404)

        objs.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} ChildrenStudyWorkRights permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })


class ChildrenStudyWorkRightsExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Children Study Work Rights',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        field_list = [f.strip() for f in fields.split(',')] if fields else list(field_header_map.keys())

        queryset = ChildrenStudyWorkRights.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'ChildrenStudyWorkRights'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'children_study_work_rights.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'children_study_work_rights.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


class ChildrenStudyWorkRightsImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=400)

        format_type = file.name.split(".")[-1].lower()
        duplicates = []
        skipped_rows = []
        required_headers = {"children study work rights"}  # Adjust column name
        optional_headers = {"description"}

        try:
            data = []

            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets,
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets,
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                dataset = Dataset()
                dataset.load(decoded_file, format="csv")
                for idx, row in enumerate(dataset.dict, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}. Found: {', '.join(row_lower.keys())}."
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            imported_count = 0
            for row in reversed(data):
                row_number = row.get("_row_number", "Unknown")
                name = str(row.get("children study work rights")) if row.get("children study work rights") else None
                description = str(row.get("description")).strip() if row.get("description") else ""

                if not name:
                    skipped_rows.append({"Row": row_number, "Reason": "Missing Study/Work Right name"})
                    continue

                existing = ChildrenStudyWorkRights.objects.filter(name__iexact=name).first()
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({"Row": row_number, "Children Study Work Rights": name, "Reason": "Already exists"})
                        continue
                    else:
                        existing.description = description
                        existing.is_deleted = False
                        existing.save()
                        imported_count += 1
                else:
                    ChildrenStudyWorkRights.objects.create(name=name, description=description, is_deleted=False)
                    imported_count += 1

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": duplicates,
            "skipped_rows": skipped_rows
        }, status=200)






class SpouseVisaCategoryListAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        search = request.GET.get('search', '').strip()
        visamain_uuid = request.GET.get('visamain_uuid')
        ordering = request.GET.get('ordering', '-created_at')  # default ordering
        page = int(request.GET.get('page', 1))
        page_size = int(request.GET.get('page_size', 20))

        qs = SpouseVisaCategory.objects.filter(is_deleted=False)

        # ---------------------- FILTER BY VISA MAIN ---------------------- #
        if visamain_uuid:
            qs = qs.filter(visamain__uuid=visamain_uuid)

        # ---------------------- SEARCH ---------------------- #
        if search:
            qs = qs.filter(
                Q(visamain__name__icontains=search) |
                Q(description__icontains=search)
            )

        # ---------------------- ORDERING ---------------------- #
        allowed_order_fields = ['created_at', 'updated_at', 'visamain__name']
        if ordering.replace('-', '') in allowed_order_fields:
            qs = qs.order_by(ordering)
        else:
            qs = qs.order_by('-created_at')

        # ---------------------- PAGINATION ---------------------- #
        total = qs.count()
        start = (page - 1) * page_size
        end = start + page_size
        qs = qs[start:end]

        serializer = SpouseVisaCategorySerializer(qs, many=True)
        return Response({
            "statusCode": 200,
            "status": True,
            "total": total,
            "page": page,
            "page_size": page_size,
            "data": serializer.data
        })
    

class SpouseVisaCategoryCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        data = request.data.copy()

        # ---------------------- FK VALIDATION ---------------------- #
        visamain_uuid = data.get("visamain_uuid")
        if not visamain_uuid:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Mandatory field missing: visamain_uuid"
            }, status=400)

        try:
            visamain_obj = VisaMain.objects.get(uuid=visamain_uuid)
        except VisaMain.DoesNotExist:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Invalid visamain UUID."
            }, status=400)

        # ---------------------- DUPLICATE CHECK ---------------------- #
        existing = SpouseVisaCategory.objects.filter(
            visamain=visamain_obj,
            is_deleted=False
        ).first()

        # If the rule is → ONE category for each visamain  
        # If multiple allowed, remove this block.
        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Category already exists for this VisaMain."
            }, status=400)

        serializer = SpouseVisaCategorySerializer(data=data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Spouse Visa Category created successfully",
                "data": serializer.data
            })

        err_msg = " ".join([msg for msgs in serializer.errors.values() for msg in msgs])
        return Response({"statusCode": 400, "status": False, "message": err_msg}, status=400)


# ---------------------- RETRIEVE ---------------------- #
class SpouseVisaCategoryRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = SpouseVisaCategory.objects.get(uuid=uuid, is_deleted=False)
        except SpouseVisaCategory.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "Not found"}, status=404)

        serializer = SpouseVisaCategorySerializer(obj)
        return Response({"statusCode": 200, "status": True, "data": serializer.data})


# ---------------------- UPDATE ---------------------- #
class SpouseVisaCategoryUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = SpouseVisaCategory.objects.get(uuid=uuid, is_deleted=False)
        except SpouseVisaCategory.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "Not found"}, status=404)

        serializer = SpouseVisaCategorySerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Updated successfully",
                "data": serializer.data
            })

        err_msg = " ".join([msg for msgs in serializer.errors.values() for msg in msgs])
        return Response({"statusCode": 400, "status": False, "message": err_msg}, status=400)


# ---------------------- DELETE ---------------------- #
class SpouseVisaCategoryDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request):
        ids = request.data.get('id')

        if not ids:
            return Response({"status": False, "message": "Provide 'id' field"}, status=400)

        if ids == "all":
            qs = SpouseVisaCategory.objects.filter(is_deleted=False)
            count = qs.count()
            qs.update(is_deleted=True)
            return Response({"status": True, "message": f"All {count} records deleted"})

        if not isinstance(ids, list):
            return Response({"status": False, "message": "Send list of UUIDs"}, status=400)

        valid, invalid = [], []
        for u in ids:
            try:
                valid.append(UUID(u))
            except:
                invalid.append(u)

        qs = SpouseVisaCategory.objects.filter(uuid__in=valid, is_deleted=False)
        count = qs.count()
        qs.update(is_deleted=True)

        return Response({
            "status": True,
            "message": f"{count} record(s) deleted",
            "invalid_uuids": invalid if invalid else None
        })


# ---------------------- EXPORT ---------------------- #
class SpouseVisaCategoryExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get("format", "xlsx").lower()
        fields = request.GET.get("fields")
        uuids = request.GET.get("uuids", "")
        uuids = [u for u in uuids.split(",") if u]

        field_header = {
            "uuid": "UUID",
            "visamain_name": "Spouse Visa Category",
            "description": "Description",
            "is_deleted": "Deleted",
            "created_at": "Created On",
            "updated_at": "Updated On",
        }

        field_list = [f.strip() for f in fields.split(',')] if fields else list(field_header.keys())

        dataset = Dataset()
        dataset.headers = [field_header.get(f, f) for f in field_list]

        qs = SpouseVisaCategory.objects.filter(is_deleted=False)
        if uuids:
            qs = qs.filter(uuid__in=uuids)

        qs = qs.order_by('-created_at')

        for obj in qs:
            row = []
            for f in field_list:
                if f == "visamain_name":
                    val = obj.visamain.name if obj.visamain else ""
                else:
                    val = getattr(obj, f, "")

                if f in ["created_at", "updated_at"] and val:
                    val = timezone.localtime(val, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                if isinstance(val, bool):
                    val = int(val)

                row.append(val)

            dataset.append(row)

        if format_type == "csv":
            file_data = dataset.export("csv")
            response = HttpResponse(file_data, content_type="text/csv")
            response["Content-Disposition"] = 'attachment; filename="spouse_visa_category.csv"'
        else:
            file_data = io.BytesIO(dataset.export("xlsx"))
            response = HttpResponse(
                file_data.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            response["Content-Disposition"] = 'attachment; filename="spouse_visa_category.xlsx"'

        return response


# ---------------------- IMPORT ---------------------- #
class SpouseVisaCategoryImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"error": "No file uploaded"}, status=400)

        format_type = file.name.split(".")[-1].lower()

        required_headers = {"spouse visa category"}
        optional_headers = {"description"}

        parsed_data = []

        try:
            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                if sheet_name not in wb.sheetnames:
                    return Response({"error": "Invalid sheet_name"}, status=400)

                ws = wb[sheet_name]
                headers = [str(c.value).strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(headers):
                    return Response({"error": f"Missing required headers: {required_headers}"}, status=400)

                for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    parsed_data.append((index, dict(zip(headers, row))))

            elif format_type == "csv":
                dataset = Dataset()
                dataset.load(file.read().decode("utf-8"), format="csv")

                for idx, row in enumerate(dataset.dict, start=2):
                    r = {k.strip().lower(): v for k, v in row.items()}
                    if not required_headers.issubset(r.keys()):
                        return Response({"error": f"Missing required headers: {required_headers}"})
                    parsed_data.append((idx, r))

            else:
                return Response({"error": "Only xlsx/csv supported"}, status=400)

            duplicate = []
            skipped = []
            imported = 0

            # ---------------------- ROW PROCESSING ---------------------- #
            for row_num, row in parsed_data:
                visa_name = str(row.get("spouse visa category")).strip()
                desc = row.get("description", "")

                if not visa_name:
                    skipped.append({"row": row_num, "Reason": "Mandatory field missing"})
                    continue

                try:
                    visa_obj = VisaMain.objects.get(name__iexact=visa_name)
                except:
                    skipped.append({"row": row_num, "Reason": "Visa Main not found"})
                    continue

                # Duplicate check
                existing = SpouseVisaCategory.objects.filter(
                    visamain=visa_obj
                ).first()

                if existing:
                    if not existing.is_deleted:
                        duplicate.append({
                            "row": row_num,
                            "Spouse Visa Category": visa_name,
                            "Reason": "Duplicate entry"
                        })
                        continue
                    else:
                        existing.description = desc
                        existing.is_deleted = False
                        existing.save()
                        imported += 1
                else:
                    SpouseVisaCategory.objects.create(
                        visamain=visa_obj,
                        description=desc,
                        is_deleted=False,
                    )
                    imported += 1

        except Exception as e:
            return Response({"error": str(e)}, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "imported_count": imported,
            "duplicates": duplicate,
            "skipped_rows": skipped,
            "message": "Import successfully completed"
        })