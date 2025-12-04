from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from rest_framework.permissions import IsAuthenticated, IsAdminUser
from django.db.models import Q
from django.http import HttpResponse
from tablib import Dataset
from tablib.formats import registry
import openpyxl, datetime
from uuid import UUID
import io, csv, datetime, uuid, openpyxl
from .models import *
from .serializers import *
import pytz
from django.utils import timezone
from .pagination import *
CSV = registry.get_format('csv')
XLSX = registry.get_format('xlsx')

india_tz = pytz.timezone('Asia/Kolkata')

# ------------------ LIST ------------------
class RepresentingCountryListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['country__name', 'continent', 'created_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = RepresentingCountry.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(
                Q(country__name__istartswith=search) |
                Q(continent__istartswith=search) |
                Q(full_name__istartswith=search) |
                Q(short_name__istartswith=search)
            )

        queryset = queryset.order_by(sort_by)

        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = RepresentingCountrySerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ------------------ CREATE ------------------
class RepresentingCountryCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        serializer = RepresentingCountrySerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Representing Country created successfully.",
                "data": serializer.data
            }, status=status.HTTP_200_OK)
        return Response({
            "statusCode": 400,
            "status": False,
            "message": serializer.errors
        }, status=status.HTTP_400_BAD_REQUEST)


# ------------------ RETRIEVE ------------------
class RepresentingCountryRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            country = RepresentingCountry.objects.get(uuid=uuid, is_deleted=False)
        except RepresentingCountry.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "Representing Country not found"
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = RepresentingCountrySerializer(country)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "Representing Country retrieved successfully",
            "data": serializer.data
        })


# ------------------ UPDATE ------------------
class RepresentingCountryUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            country = RepresentingCountry.objects.get(uuid=uuid, is_deleted=False)
        except RepresentingCountry.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "Representing Country not found"
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = RepresentingCountrySerializer(country, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "Representing Country updated successfully",
                "data": serializer.data
            })
        return Response({
            "statusCode": 400,
            "status": False,
            "message": serializer.errors
        }, status=status.HTTP_400_BAD_REQUEST)


# ------------------ DELETE ------------------
class RepresentingCountryDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request):
        ids = request.data.get('id', None)
        if not ids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Please provide 'id' field (UUID list or 'all')."
            }, status=status.HTTP_400_BAD_REQUEST)

        if ids == "all":
            countries = RepresentingCountry.objects.filter(is_deleted=False)
            count = countries.count()
            countries.update(is_deleted=True)
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} country records deleted successfully."
            })

        if not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Provide list of UUIDs or 'all'."
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids, invalid_uuids = [], []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        countries = RepresentingCountry.objects.filter(uuid__in=valid_uuids, is_deleted=False)
        count = countries.count()
        countries.update(is_deleted=True)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} country record(s) deleted successfully.",
            "data": {"invalid_uuids": invalid_uuids}
        })


# ------------------ EXPORT ------------------
class RepresentingCountryExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')  # comma-separated fields
        uuids_param = request.GET.get('uuids', '')  # comma-separated UUIDs
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        # --- Field to header mapping ---
        field_header_map = {
            'uuid': 'UUID',
            'country': 'Country Name',
            'continent': 'Continent',
            'short_name': 'Short Name',
            'full_name': 'Full Name',
            'official_name': 'Official Name',
            'capital_city': 'Capital City',
            'population': 'Population',
            'status': 'Status',
            'is_active': 'Active',
            'is_deleted': 'Deleted',
            'created_at': 'Created On',
            'updated_at': 'Modified On'
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = RepresentingCountry.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'Representing Countries'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field == 'continent' and obj.continent:
                    value = obj.continent
                elif isinstance(value, datetime.datetime):
                    value = timezone.localtime(value).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv; charset=utf-8'
            file_name = 'representing_countries.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'representing_countries.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


# ------------------ IMPORT ------------------
class RepresentingCountryImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get('file')
        sheet_name = request.data.get('sheet_name')

        if not file:
            return Response({'error': 'No file uploaded'}, status=400)

        format_type = file.name.split('.')[-1].lower()

        required_headers = {'full name', 'continent'}
        optional_headers = {'short name', 'official name', 'capital city', 'population', 'status'}

        all_headers = required_headers.union(optional_headers)

        data = []
        duplicates = []
        skipped_rows = []

        try:
            # ---------------- XLSX ----------------
            if format_type == 'xlsx':
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        'error': 'Please provide sheet_name',
                        'available_sheets': available_sheets
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        'error': f'Sheet "{sheet_name}" not found',
                        'available_sheets': available_sheets
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [
                    str(c.value).strip().lower() if c.value else ''
                    for c in next(ws.iter_rows(min_row=1, max_row=1))
                ]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = row_no
                    data.append(row_dict)

            # ---------------- CSV ----------------
            elif format_type == 'csv':
                decoded_file = file.read().decode('utf-8')
                reader = csv.DictReader(io.StringIO(decoded_file))

                for row_no, row in enumerate(reader, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = row_no

                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {required_headers}"
                        }, status=400)

                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv"
                }, status=400)

            # ---------------- PROCESSING ----------------
            imported_count = 0
            to_create = []
            seen_in_file = set()

            continents_map = {c.name.lower(): c for c in Continents.objects.all()}

            existing_map = {
                c.full_name.lower(): c
                for c in RepresentingCountry.objects.all()
            }

            for row in reversed(data):

                row_no = row.get("_row_number", "Unknown")

                # Extract all fields
                full_name = str(row.get("full name") or "").strip()
                short_name = str(row.get("short name") or "").strip()
                official_name = str(row.get("official name") or "").strip()
                capital_city = str(row.get("capital city") or "").strip()
                population = str(row.get("population") or "").strip()
                status = str(row.get("status") or "").strip()
                continent_name = str(row.get("continent") or "").strip()

                lower_full = full_name.lower()

                # ---------------- Missing Required Field ----------------
                if not full_name:
                    skipped_rows.append({
                        "Row": row_no,
                        "Full Name": full_name,
                        "Short Name": short_name,
                        "Official Name": official_name,
                        "Capital City": capital_city,
                        "Population": population,
                        "Status": status,
                        "Continent": continent_name,
                        "Reason": "Missing required field: full name"
                    })
                    continue

                # ---------------- Validate Continent ----------------
                continent_obj = None
                if continent_name:
                    continent_obj = continents_map.get(continent_name.lower())
                    if not continent_obj:
                        skipped_rows.append({
                            "Row": row_no,
                            "Full Name": full_name,
                            "Short Name": short_name,
                            "Official Name": official_name,
                            "Capital City": capital_city,
                            "Population": population,
                            "Status": status,
                            "Continent": continent_name,
                            "Reason": "Invalid continent name"
                        })
                        continue

                # ---------------- Duplicate Inside Import File ----------------
                if lower_full in seen_in_file:
                    duplicates.append({
                        "Row": row_no,
                        "Full Name": full_name,
                        "Short Name": short_name,
                        "Official Name": official_name,
                        "Capital City": capital_city,
                        "Population": population,
                        "Status": status,
                        "Continent": continent_name,
                        "Reason": "Duplicate in uploaded file"
                    })
                    continue

                seen_in_file.add(lower_full)

                # ---------------- Already in Database ----------------
                existing = existing_map.get(lower_full)

                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "Row": row_no,
                            "Full Name": full_name,
                            "Short Name": short_name,
                            "Official Name": official_name,
                            "Capital City": capital_city,
                            "Population": population,
                            "Status": status,
                            "Continent": continent_name,
                            "Reason": "Already exists in database"
                        })
                        continue

                    # Reactivate Deleted Record
                    existing.continent = continent_obj
                    existing.short_name = short_name
                    existing.official_name = official_name
                    existing.capital_city = capital_city
                    existing.population = population
                    existing.status = status
                    existing.is_deleted = False
                    existing.save()

                    imported_count += 1
                    continue

                # ---------------- Add New Row for Bulk Create ----------------
                to_create.append(
                    RepresentingCountry(
                        full_name=full_name,
                        short_name=short_name,
                        official_name=official_name,
                        capital_city=capital_city,
                        population=population,
                        status=status,
                        continent=continent_obj,
                        is_deleted=False
                    )
                )

            # ---------------- BULK CREATE ----------------
            if to_create:
                batch_size = 300
                for i in range(0, len(to_create), batch_size):
                    RepresentingCountry.objects.bulk_create(to_create[i:i + batch_size])

                imported_count += len(to_create)

        except Exception as e:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": str(e)
            }, status=400)

        # ---------------- FINAL RESPONSE ----------------
        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": list(reversed(duplicates)),
            "skipped_rows": list(reversed(skipped_rows)),
        }, status=200)




# ------------------ VisaMain ------------------
class VisaMainListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'created_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'
        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = VisaMain.objects.filter(is_deleted=False)
        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = VisaMainSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ------------------ CREATE ------------------
class VisaMainCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get('name', '').strip()
        if VisaMain.objects.filter(name__iexact=name, is_deleted=False).exists():
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "VisaMain with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = VisaMainSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "VisaMain created successfully.",
                "data": serializer.data
            })
        return Response({"statusCode": 400, "status": False, "message": serializer.errors}, status=400)


# ------------------ RETRIEVE ------------------
class VisaMainRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = VisaMain.objects.get(uuid=uuid, is_deleted=False)
        except VisaMain.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "VisaMain not found"}, status=404)

        serializer = VisaMainSerializer(obj)
        return Response({"statusCode": 200, "status": True, "message": "Retrieved successfully", "data": serializer.data})


# ------------------ UPDATE ------------------
class VisaMainUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = VisaMain.objects.get(uuid=uuid, is_deleted=False)
        except VisaMain.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "VisaMain not found"}, status=404)

        serializer = VisaMainSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({"statusCode": 200, "status": True, "message": "Updated successfully", "data": serializer.data})
        return Response({"statusCode": 400, "status": False, "message": serializer.errors}, status=400)


# ------------------ DELETE ------------------
class VisaMainDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request):
        ids = request.data.get('id')
        if not ids:
            return Response({"statusCode": 400, "status": False, "message": "Provide 'id' field (list of UUIDs or 'all')."}, status=400)

        if ids == "all":
            objs = VisaMain.objects.filter(is_deleted=False)
            count = objs.count()
            objs.update(is_deleted=True)
            return Response({"statusCode": 200, "status": True, "message": f"All {count} VisaMain(s) deleted."})

        if not isinstance(ids, list):
            return Response({"statusCode": 400, "status": False, "message": "Provide a list of UUIDs."}, status=400)

        valid_uuids, invalid_uuids = [], []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        objs = VisaMain.objects.filter(uuid__in=valid_uuids, is_deleted=False)
        count = objs.count()
        objs.update(is_deleted=True)

        return Response({"statusCode": 200, "status": True, "message": f"{count} VisaMain(s) deleted.", "data": {"invalid_uuids": invalid_uuids}})


# ------------------ EXPORT ------------------
class VisaMainExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        # Field mapping
        field_header_map = {
            'uuid': 'UUID',
            'name': 'Visa Main',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'created_at': 'Created On',
            'updated_at': 'Modified On',
        }

        field_list = [f.strip() for f in fields.split(',')] if fields else list(field_header_map.keys())

        queryset = VisaMain.objects.filter(uuid__in=uuids) if uuids else VisaMain.objects.all()
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'VisaMain'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv; charset=utf-8'
            file_name = 'visamain.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'visamain.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


# ------------------ IMPORT ------------------
class VisaMainImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get('file')
        sheet_name = request.data.get('sheet_name')

        if not file:
            return Response({"error": "No file uploaded"}, status=400)

        format_type = file.name.split('.')[-1].lower()
        required_headers = {"visa main"}
        optional_headers = {"description"}
        all_headers = required_headers.union(optional_headers)

        data = []
        duplicates = []
        skipped_rows = []
        to_create = []

        try:
            # ---------- XLSX Handling ----------
            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "error": "Please provide sheet_name",
                        "available_sheets": available_sheets
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "error": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets
                    }, status=400)

                ws = wb[sheet_name]

                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = row_no
                    data.append(row_dict)

            # ---------- CSV Handling ----------
            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                reader = csv.DictReader(io.StringIO(decoded_file))

                for row_no, row in enumerate(reader, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = row_no

                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {required_headers}, Found: {set(row_lower.keys())}"
                        }, status=400)

                    data.append(row_lower)
            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv",
                }, status=400)

            # ---------- Processing Rows ----------
            imported_count = 0
            seen_in_file = set()

            existing_map = {v.name.lower(): v for v in VisaMain.objects.all()}

            for row in reversed(data):
                row_no = row.get("_row_number", "Unknown")
                name = str(row.get("visa main") or "").strip()
                description = str(row.get("description") or "").strip()

                if not name:
                    skipped_rows.append({
                        "Row": row_no,
                        "Visa Main": name,
                        "Description": description,
                        "Reason": "Missing required field: visa main"
                    })
                    continue

                lower_name = name.lower()

                # ---------- Duplicate in uploaded file ----------
                if lower_name in seen_in_file:
                    duplicates.append({
                        "Row": row_no,
                        "Visa Main": name,
                        "Description": description,
                        "Reason": "Duplicate in uploaded file"
                    })
                    continue
                seen_in_file.add(lower_name)

                # ---------- Existing record ----------
                existing = existing_map.get(lower_name)
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "Row": row_no,
                            "Visa Main": name,
                            "Description": description,
                            "Reason": "Already exists in database"
                        })
                        continue
                    # Reactivate deleted record
                    existing.description = description
                    existing.is_deleted = False
                    existing.save()
                    imported_count += 1
                    continue

                # ---------- Add to bulk_create list ----------
                to_create.append(VisaMain(
                    name=name,
                    description=description,
                    is_deleted=False
                ))

            # ---------- Bulk create ----------
            if to_create:
                batch_size = 500
                for i in range(0, len(to_create), batch_size):
                    VisaMain.objects.bulk_create(to_create[i:i + batch_size])
                imported_count += len(to_create)

        except Exception as e:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": str(e)
            }, status=400)

        # ---------- Response ----------
        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": list(reversed(duplicates)),
            "skipped_rows": list(reversed(skipped_rows)),
        }, status=200)



#--------------------visamajor-----------------
class VisaMajorListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')
        allowed_sort_fields = ['name', 'created_at']

        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'
        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = VisaMajor.objects.filter(is_deleted=False)
        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = VisaMajorSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ------------------ CREATE ------------------
class VisaMajorCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        visamain_id = request.data.get('visamain')
        name = request.data.get('name', '').strip()

        if not visamain_id:
            return Response({"statusCode": 400, "status": False, "message": "visamain is required"}, status=400)

        if VisaMajor.objects.filter(visamain_id=visamain_id, name__iexact=name, is_deleted=False).exists():
            return Response({"statusCode": 400, "status": False, "message": "VisaMajor with this name already exists for this VisaMain"}, status=400)

        serializer = VisaMajorSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({"statusCode": 200, "status": True, "message": "VisaMajor created successfully", "data": serializer.data})
        return Response({"statusCode": 400, "status": False, "message": serializer.errors}, status=400)


# ------------------ RETRIEVE ------------------
class VisaMajorRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = VisaMajor.objects.get(uuid=uuid, is_deleted=False)
        except VisaMajor.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "VisaMajor not found"}, status=404)

        serializer = VisaMajorSerializer(obj)
        return Response({"statusCode": 200, "status": True, "message": "Retrieved successfully", "data": serializer.data})


# ------------------ UPDATE ------------------
class VisaMajorUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = VisaMajor.objects.get(uuid=uuid, is_deleted=False)
        except VisaMajor.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "VisaMajor not found"}, status=404)

        visamain_id = request.data.get('visamain')
        name = request.data.get('name', '').strip()

        if VisaMajor.objects.filter(visamain_id=visamain_id, name__iexact=name, is_deleted=False).exclude(uuid=uuid).exists():
            return Response({"statusCode": 400, "status": False, "message": "VisaMajor with this name already exists for this VisaMain"}, status=400)

        serializer = VisaMajorSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({"statusCode": 200, "status": True, "message": "Updated successfully", "data": serializer.data})
        return Response({"statusCode": 400, "status": False, "message": serializer.errors}, status=400)


# ------------------ DELETE ------------------
class VisaMajorDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request):
        ids = request.data.get('id')
        if not ids:
            return Response({"statusCode": 400, "status": False, "message": "Provide 'id' field (list of UUIDs or 'all')."}, status=400)

        if ids == "all":
            objs = VisaMajor.objects.filter(is_deleted=False)
            count = objs.count()
            objs.update(is_deleted=True)
            return Response({"statusCode": 200, "status": True, "message": f"All {count} VisaMajor(s) deleted."})

        if not isinstance(ids, list):
            return Response({"statusCode": 400, "status": False, "message": "Provide a list of UUIDs."}, status=400)

        valid_uuids, invalid_uuids = [], []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        objs = VisaMajor.objects.filter(uuid__in=valid_uuids, is_deleted=False)
        count = objs.count()
        objs.update(is_deleted=True)

        return Response({"statusCode": 200, "status": True, "message": f"{count} VisaMajor(s) deleted.", "data": {"invalid_uuids": invalid_uuids}})


# ------------------ EXPORT ------------------
class VisaMajorExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        # --- Field to header mapping ---
        field_header_map = {
            'uuid': 'UUID',
            'country': 'Country',
            'visamain': 'Visa Main',
            'visamain_name': 'Visa Main Name',
            'name': 'Visa Major Name',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'created_at': 'Created On',
            'updated_at': 'Modified On'
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = VisaMajor.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'VisaMajor'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')

                # handle related fields
                if field == 'country' and obj.country:
                    value = obj.country.name
                elif field == 'visamain_name' and obj.visamain:
                    value = obj.visamain.name
                elif field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)

                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv; charset=utf-8'
            file_name = 'visamajor.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'visamajor.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


# ------------------ IMPORT ------------------
class VisaMajorImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get('file')
        sheet_name = request.data.get('sheet_name')

        if not file:
            return Response({'error': 'No file uploaded'}, status=400)

        format_type = file.name.split('.')[-1].lower()
        required_headers = {'country', 'visamain', 'name'}
        optional_headers = {'description'}
        all_headers = required_headers.union(optional_headers)

        data = []
        duplicates = []
        skipped_rows = []
        to_create = []

        try:
            # ---------------- XLSX Handling ----------------
            if format_type == 'xlsx':
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({'statusCode': 400, 'status': False, 'error': 'Please provide sheet_name', 'available_sheets': available_sheets}, status=400)
                if sheet_name not in available_sheets:
                    return Response({'statusCode': 400, 'status': False, 'error': f'Sheet "{sheet_name}" not found', 'available_sheets': available_sheets}, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({'statusCode': 400, 'status': False, 'message': f'Sheet "{sheet_name}" is empty'}, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else '' for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                if not required_headers.issubset(set(headers)):
                    return Response({'statusCode': 400, 'status': False, 'message': f'Missing required headers: {required_headers}. Found: {set(headers)}'}, status=400)

                for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = row_no
                    data.append(row_dict)

            # ---------------- CSV Handling ----------------
            elif format_type == 'csv':
                decoded_file = file.read().decode('utf-8')
                reader = csv.DictReader(io.StringIO(decoded_file))

                for row_no, row in enumerate(reader, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = row_no

                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({'statusCode': 400, 'status': False, 'message': f'Missing required headers: {required_headers}. Found: {set(row_lower.keys())}'}, status=400)

                    data.append(row_lower)

            else:
                return Response({'statusCode': 400, 'status': False, 'message': 'Unsupported file format. Use .xlsx or .csv'}, status=400)

            # ---------------- Processing Rows ----------------
            imported_count = 0
            seen_in_file = set()

            # Fetch related objects in bulk to reduce queries
            countries_map = {c.full_name.lower(): c for c in RepresentingCountry.objects.all()}
            visamain_map = {v.name.lower(): v for v in VisaMain.objects.all()}
            existing_map = {
                (v.country.id, v.visamain.id, v.name.lower()): v
                for v in VisaMajor.objects.select_related('country', 'visamain').all()
            }

            for row in reversed(data):
                row_no = row.get("_row_number", "Unknown")
                country_name = str(row.get("country") or "").strip()
                visamain_name = str(row.get("visamain") or "").strip()
                name = str(row.get("name") or "").strip()
                description = str(row.get("description") or "").strip()

                key_in_file = (country_name.lower(), visamain_name.lower(), name.lower())

                if not country_name or not visamain_name or not name:
                    skipped_rows.append({
                        "Row": row_no,
                        "Country": country_name,
                        "VisaMain": visamain_name,
                        "Name": name,
                        "Description": description,
                        "Reason": "Missing required fields"
                    })
                    continue

                # Duplicate in uploaded file
                if key_in_file in seen_in_file:
                    duplicates.append({
                        "Row": row_no,
                        "Country": country_name,
                        "VisaMain": visamain_name,
                        "Name": name,
                        "Description": description,
                        "Reason": "Duplicate in uploaded file"
                    })
                    continue
                seen_in_file.add(key_in_file)

                # Validate related objects
                country_obj = countries_map.get(country_name.lower())
                visamain_obj = visamain_map.get(visamain_name.lower())
                if not country_obj or not visamain_obj:
                    skipped_rows.append({
                        "Row": row_no,
                        "Country": country_name,
                        "VisaMain": visamain_name,
                        "Name": name,
                        "Description": description,
                        "Reason": "Invalid country or VisaMain"
                    })
                    continue

                # Check existing in DB
                existing = existing_map.get((country_obj.id, visamain_obj.id, name.lower()))
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "Row": row_no,
                            "Country": country_name,
                            "VisaMain": visamain_name,
                            "Name": name,
                            "Description": description,
                            "Reason": "Already exists in database"
                        })
                        continue
                    # Reactivate deleted record
                    existing.description = description
                    existing.is_deleted = False
                    existing.save()
                    imported_count += 1
                    continue

                # Prepare for bulk create
                to_create.append(VisaMajor(
                    country=country_obj,
                    visamain=visamain_obj,
                    name=name,
                    description=description,
                    is_deleted=False
                ))

            # ---------------- Bulk Create ----------------
            if to_create:
                batch_size = 500
                for i in range(0, len(to_create), batch_size):
                    VisaMajor.objects.bulk_create(to_create[i:i + batch_size])
                imported_count += len(to_create)

        except Exception as e:
            return Response({'statusCode': 400, 'status': False, 'message': str(e)}, status=400)

        # ---------------- Response ----------------
        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": list(reversed(duplicates)),
            "skipped_rows": list(reversed(skipped_rows)),
        }, status=200)

#-------------------------VisaName----------------
class VisaNameListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')
        allowed_sort_fields = ['full_name', 'created_at']

        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'
        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = VisaName.objects.filter(is_deleted=False)
        if search:
            queryset = queryset.filter(
                Q(full_name__istartswith=search) |
                Q(short_name__istartswith=search) |
                Q(description__istartswith=search)
            )

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = VisaNameSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ------------------ CREATE ------------------
class VisaNameCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        serializer = VisaNameSerializer(data=request.data)

        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "VisaName created successfully",
                "data": serializer.data
            }, status=200)

        return Response({
            "statusCode": 400,
            "status": False,
            "message": serializer.errors
        }, status=400)


# ------------------ RETRIEVE ------------------
class VisaNameRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = VisaName.objects.get(uuid=uuid, is_deleted=False)
        except VisaName.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "VisaName not found"}, status=404)

        serializer = VisaNameSerializer(obj)
        return Response({"statusCode": 200, "status": True, "message": "Retrieved successfully", "data": serializer.data})


# ------------------ UPDATE ------------------
class VisaNameUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = VisaName.objects.get(uuid=uuid, is_deleted=False)
        except VisaName.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "VisaName not found"}, status=404)

        country_id = request.data.get('country')
        visamain_id = request.data.get('visamain')
        visamajor_id = request.data.get('visamajor')
        full_name = request.data.get('full_name', '').strip()

        if VisaName.objects.filter(
            country_id=country_id,
            visamain_id=visamain_id,
            visamajor_id=visamajor_id,
            full_name__iexact=full_name,
            is_deleted=False
        ).exclude(uuid=uuid).exists():
            return Response({"statusCode": 400, "status": False, "message": "VisaName with this combination already exists"}, status=400)

        serializer = VisaNameSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({"statusCode": 200, "status": True, "message": "Updated successfully", "data": serializer.data})
        return Response({"statusCode": 400, "status": False, "message": serializer.errors}, status=400)


# ------------------ DELETE ------------------
class VisaNameDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request):
        ids = request.data.get('id')
        if not ids:
            return Response({"statusCode": 400, "status": False, "message": "Provide 'id' field (list of UUIDs or 'all')."}, status=400)

        if ids == "all":
            objs = VisaName.objects.filter(is_deleted=False)
            count = objs.count()
            objs.update(is_deleted=True)
            return Response({"statusCode": 200, "status": True, "message": f"All {count} VisaName(s) deleted."})

        if not isinstance(ids, list):
            return Response({"statusCode": 400, "status": False, "message": "Provide a list of UUIDs."}, status=400)

        valid_uuids, invalid_uuids = [], []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        objs = VisaName.objects.filter(uuid__in=valid_uuids, is_deleted=False)
        count = objs.count()
        objs.update(is_deleted=True)

        return Response({"statusCode": 200, "status": True, "message": f"{count} VisaName(s) deleted.", "data": {"invalid_uuids": invalid_uuids}})


# ------------------ EXPORT ------------------
class VisaNameExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'csv').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]
        field_list = [f.strip() for f in fields.split(',')] if fields else [
            'uuid', 'country', 'country_name', 'visamain', 'visamain_name',
            'visamajor', 'visamajor_name', 'full_name', 'short_name', 'description', 'is_deleted', 'created_at', 'updated_at'
        ]

        queryset = VisaName.objects.filter(uuid__in=uuids) if uuids else VisaName.objects.all()
        dataset = Dataset()
        dataset.headers = field_list

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if isinstance(value, datetime.datetime):
                    value = value.strftime("%Y-%m-%d %H:%M:%S")
                if isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'xlsx':
            data = XLSX().export_data(dataset)
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            filename = 'visaname.xlsx'
        else:
            data = CSV().export_data(dataset)
            content_type = 'text/csv; charset=utf-8'
            filename = 'visaname.csv'

        response = HttpResponse(data, content_type=content_type)
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


# ------------------ IMPORT ------------------
class VisaNameImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No file uploaded"
            }, status=400)

        format_type = file.name.split(".")[-1].lower()
        required_headers = {'country', 'visa main category', 'visa major category', 'visa   name'}
        optional_headers = {'visa short name', 'description'}

        data = []
        duplicates = []
        skipped_rows = []
        to_create = []

        try:
            # ---------------- XLSX Handling ----------------
            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": "Please provide sheet_name",
                        "available_sheets": available_sheets
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty'
                    }, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            # ---------------- CSV Handling ----------------
            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                reader = csv.DictReader(io.StringIO(decoded_file))
                for idx, row in enumerate(reader, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {required_headers}, Found: {set(row_lower.keys())}"
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv"
                }, status=400)

            imported_count = 0
            seen_in_file = set()

            # Preload existing records
            existing_map = {}
            for vn in VisaName.objects.all():
                key = (
                    vn.country_id,
                    vn.visamain_id,
                    vn.visamajor_id,
                    vn.full_name.lower()
                )
                existing_map[key] = vn

            # ---------------- Process Rows ----------------
            for row in reversed(data):
                row_no = row.get("_row_number", "Unknown")
                country_name = str(row.get('country') or '').strip()
                visamain_name = str(row.get('visa main category') or '').strip()
                visamajor_name = str(row.get('visa major category') or '').strip()
                full_name = str(row.get('visa   name') or '').strip()
                short_name = str(row.get('visa short name') or '').strip()
                description = str(row.get('description') or '').strip()

                # Validate required fields
                if not country_name or not visamain_name or not visamajor_name or not full_name:
                    skipped_rows.append({
                        "Row": row_no,
                        "Country": country_name,
                        "VisaMain": visamain_name,
                        "VisaMajor": visamajor_name,
                        "VisaName": full_name,
                        "Reason": "Missing required fields"
                    })
                    continue

                # Fetch related objects
                country_obj = RepresentingCountry.objects.filter(full_name__iexact=country_name).first()
                visamain_obj = VisaMain.objects.filter(name__iexact=visamain_name).first()
                visamajor_obj = VisaMajor.objects.filter(name__iexact=visamajor_name).first()

                if not country_obj or not visamain_obj or not visamajor_obj:
                    skipped_rows.append({
                        "Row": row_no,
                        "Country": country_name,
                        "VisaMain": visamain_name,
                        "VisaMajor": visamajor_name,
                        "VisaName": full_name,
                        "Reason": "Invalid country / visamain / visamajor"
                    })
                    continue

                key = (country_obj.id, visamain_obj.id, visamajor_obj.id, full_name.lower())

                # Duplicate in file
                if key in seen_in_file:
                    duplicates.append({
                        "Row": row_no,
                        "Country": country_name,
                        "VisaMain": visamain_name,
                        "VisaMajor": visamajor_name,
                        "VisaName": full_name,
                        "Reason": "Duplicate in uploaded file"
                    })
                    continue
                seen_in_file.add(key)

                existing = existing_map.get(key)
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "Row": row_no,
                            "Country": country_name,
                            "VisaMain": visamain_name,
                            "VisaMajor": visamajor_name,
                            "VisaName": full_name,
                            "Reason": "Already exists in database"
                        })
                        continue
                    # Reactivate deleted record
                    existing.short_name = short_name
                    existing.description = description
                    existing.is_deleted = False
                    existing.save()
                    imported_count += 1
                    continue

                # Prepare for bulk create
                to_create.append(VisaName(
                    country=country_obj,
                    visamain=visamain_obj,
                    visamajor=visamajor_obj,
                    full_name=full_name,
                    short_name=short_name,
                    description=description,
                    is_deleted=False
                ))

            # ---------------- Bulk Create ----------------
            if to_create:
                batch_size = 500
                for i in range(0, len(to_create), batch_size):
                    VisaName.objects.bulk_create(to_create[i:i + batch_size])
                imported_count += len(to_create)

        except Exception as e:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": str(e)
            }, status=400)

        # ---------------- SUCCESS RESPONSE -----------------
        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": list(reversed(duplicates)),
            "skipped_rows": list(reversed(skipped_rows))
        }, status=200)



#----------------applicantType--------------------
class ApplicantTypeListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')
        allowed_sort_fields = ['name', 'created_at']

        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'
        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = ApplicantType.objects.filter(is_deleted=False)
        if search:
            queryset = queryset.filter(
                Q(name__istartswith=search)
            )

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = ApplicantTypeSerializer(result_page, many=True)
        return paginator.get_paginated_response(serializer.data)


# ------------------ CREATE ------------------
class ApplicantTypeCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get('name', '').strip()
        if not name:
            return Response({"statusCode": 400, "status": False, "message": "Name is required"}, status=400)

        if ApplicantType.objects.filter(name__iexact=name, is_deleted=False).exists():
            return Response({"statusCode": 400, "status": False, "message": "ApplicantType with this name already exists"}, status=400)

        serializer = ApplicantTypeSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({"statusCode": 200, "status": True, "message": "ApplicantType created successfully", "data": serializer.data})
        return Response({"statusCode": 400, "status": False, "message": serializer.errors}, status=400)


# ------------------ RETRIEVE ------------------
class ApplicantTypeRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = ApplicantType.objects.get(uuid=uuid, is_deleted=False)
        except ApplicantType.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "ApplicantType not found"}, status=404)

        serializer = ApplicantTypeSerializer(obj)
        return Response({"statusCode": 200, "status": True, "message": "Retrieved successfully", "data": serializer.data})


# ------------------ UPDATE ------------------
class ApplicantTypeUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = ApplicantType.objects.get(uuid=uuid, is_deleted=False)
        except ApplicantType.DoesNotExist:
            return Response({"statusCode": 404, "status": False, "message": "ApplicantType not found"}, status=404)

        name = request.data.get('name', '').strip()
        if ApplicantType.objects.filter(name__iexact=name, is_deleted=False).exclude(uuid=uuid).exists():
            return Response({"statusCode": 400, "status": False, "message": "ApplicantType with this name already exists"}, status=400)

        serializer = ApplicantTypeSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({"statusCode": 200, "status": True, "message": "Updated successfully", "data": serializer.data})
        return Response({"statusCode": 400, "status": False, "message": serializer.errors}, status=400)


# ------------------ DELETE ------------------
class ApplicantTypeDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request):
        ids = request.data.get('id')
        if not ids:
            return Response({"statusCode": 400, "status": False, "message": "Provide 'id' field (list of UUIDs or 'all')."}, status=400)

        if ids == "all":
            objs = ApplicantType.objects.filter(is_deleted=False)
            count = objs.count()
            objs.update(is_deleted=True)
            return Response({"statusCode": 200, "status": True, "message": f"All {count} ApplicantType(s) deleted."})

        if not isinstance(ids, list):
            return Response({"statusCode": 400, "status": False, "message": "Provide a list of UUIDs."}, status=400)

        valid_uuids, invalid_uuids = [], []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        objs = ApplicantType.objects.filter(uuid__in=valid_uuids, is_deleted=False)
        count = objs.count()
        objs.update(is_deleted=True)

        return Response({"statusCode": 200, "status": True, "message": f"{count} ApplicantType(s) deleted.", "data": {"invalid_uuids": invalid_uuids}})


# ------------------ EXPORT ------------------
class ApplicantTypeExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'csv').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]
        field_list = [f.strip() for f in fields.split(',')] if fields else [
            'uuid', 'name', 'description', 'is_deleted', 'created_at', 'updated_at'
        ]

        queryset = ApplicantType.objects.filter(uuid__in=uuids) if uuids else ApplicantType.objects.all()
        dataset = Dataset()
        dataset.headers = field_list

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if isinstance(value, datetime.datetime):
                    value = value.strftime("%Y-%m-%d %H:%M:%S")
                if isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'xlsx':
            data = XLSX().export_data(dataset)
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            filename = 'applicanttype.xlsx'
        else:
            data = CSV().export_data(dataset)
            content_type = 'text/csv; charset=utf-8'
            filename = 'applicanttype.csv'

        response = HttpResponse(data, content_type=content_type)
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


# ------------------ IMPORT ------------------
class ApplicantTypeImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get('file')
        sheet_name = request.data.get('sheet_name')

        if not file:
            return Response({'error': 'No file uploaded'}, status=400)

        format_type = file.name.split('.')[-1].lower()
        required_headers = {'applicant type'}
        optional_headers = {'description'}
        all_headers = required_headers.union(optional_headers)

        data = []
        duplicates = []
        skipped_rows = []
        to_create = []

        try:
            # ---------------- XLSX Handling ----------------
            if format_type == 'xlsx':
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({'error': 'Please provide sheet_name', 'available_sheets': available_sheets}, status=400)
                if sheet_name not in available_sheets:
                    return Response({'error': f'Sheet "{sheet_name}" not found', 'available_sheets': available_sheets}, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({'statusCode': 400, 'status': False, 'message': f'Sheet "{sheet_name}" is empty'}, status=400)

                headers = [str(cell.value).strip().lower() if cell.value else '' for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                if not required_headers.issubset(set(headers)):
                    return Response({'statusCode': 400, 'status': False, 'message': f'Missing required headers: {required_headers}. Found: {set(headers)}'}, status=400)

                for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = row_no
                    data.append(row_dict)

            # ---------------- CSV Handling ----------------
            elif format_type == 'csv':
                decoded_file = file.read().decode('utf-8')
                reader = csv.DictReader(io.StringIO(decoded_file))
                for row_no, row in enumerate(reader, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = row_no
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({'statusCode': 400, 'status': False, 'message': f'Missing required headers: {required_headers}. Found: {set(row_lower.keys())}'}, status=400)
                    data.append(row_lower)

            else:
                return Response({'statusCode': 400, 'status': False, 'message': 'Unsupported file format. Use .xlsx or .csv'}, status=400)

            # ---------------- Process Rows ----------------
            imported_count = 0
            seen_in_file = set()

            # Fetch all existing ApplicantType to reduce DB queries
            existing_map = {a.name.lower(): a for a in ApplicantType.objects.all()}

            for row in reversed(data):
                row_no = row.get("_row_number", "Unknown")
                name = str(row.get("applicant type") or "").strip()
                description = str(row.get("description") or "").strip()

                key_in_file = name.lower()

                # Missing required field
                if not name:
                    skipped_rows.append({
                        "Row": row_no,
                        "Applicant Type": name,
                        "Description": description,
                        "Reason": "Missing required field: applicant type"
                    })
                    continue

                # Duplicate in uploaded file
                if key_in_file in seen_in_file:
                    duplicates.append({
                        "Row": row_no,
                        "Applicant Type": name,
                        "Description": description,
                        "Reason": "Duplicate in uploaded file"
                    })
                    continue
                seen_in_file.add(key_in_file)

                # Check existing in DB
                existing = existing_map.get(key_in_file)
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "Row": row_no,
                            "Applicant Type": name,
                            "Description": description,
                            "Reason": "Already exists in database"
                        })
                        continue
                    # Reactivate deleted record
                    existing.description = description
                    existing.is_deleted = False
                    existing.save()
                    imported_count += 1
                    continue

                # Prepare for bulk create
                to_create.append(ApplicantType(
                    name=name,
                    description=description,
                    is_deleted=False
                ))

            # ---------------- Bulk Create ----------------
            if to_create:
                batch_size = 500
                for i in range(0, len(to_create), batch_size):
                    ApplicantType.objects.bulk_create(to_create[i:i + batch_size])
                imported_count += len(to_create)

        except Exception as e:
            return Response({'statusCode': 400, 'status': False, 'message': str(e)}, status=400)

        # ---------------- Response ----------------
        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": list(reversed(duplicates)),
            "skipped_rows": list(reversed(skipped_rows)),
        }, status=200)


class VisaEligibilityTypeListAPIView(APIView):
    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f"-{sort_by}"

        queryset = VisaEligibilityType.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)
        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = VisaEligibilityTypeSerializer(result_page, many=True)

        return paginator.get_paginated_response(serializer.data)


class VisaEligibilityTypeCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = VisaEligibilityType.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "VisaEligibilityType with this name already exists."
            }, status=400)

        serializer = VisaEligibilityTypeSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "VisaEligibilityType created successfully",
                "data": serializer.data
            })
        else:
            errors = " ".join([msg for msgs in serializer.errors.values() for msg in msgs])
            return Response({
                "statusCode": 400,
                "status": False,
                "message": errors
            }, status=400)



class VisaEligibilityTypeRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = VisaEligibilityType.objects.get(uuid=uuid, is_deleted=False)
        except VisaEligibilityType.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "VisaEligibilityType not found",
                "data": None
            }, status=404)

        serializer = VisaEligibilityTypeSerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "VisaEligibilityType retrieved successfully",
            "data": serializer.data
        })



class VisaEligibilityTypeUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = VisaEligibilityType.objects.get(uuid=uuid, is_deleted=False)
        except VisaEligibilityType.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "VisaEligibilityType not found",
                "data": None
            }, status=404)

        serializer = VisaEligibilityTypeSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "VisaEligibilityType updated successfully",
                "data": serializer.data
            })
        else:
            errors = " ".join([msg for msgs in serializer.errors.values() for msg in msgs])
            return Response({
                "statusCode": 400,
                "status": False,
                "message": errors
            }, status=400)





class VisaEligibilityTypeDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id')

        # Delete by UUID
        if uuid:
            try:
                obj = VisaEligibilityType.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "VisaEligibilityType permanently deleted."
                }, status=204)
            except VisaEligibilityType.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "VisaEligibilityType not found."
                }, status=404)

        # Delete all
        if ids == "all":
            all_items = VisaEligibilityType.objects.all()
            count = all_items.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No VisaEligibilityType records to delete."
                }, status=404)
            all_items.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} VisaEligibilityType entries deleted."
            })

        # Multiple delete
        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Provide list of UUIDs or 'all'."
            }, status=400)

        valid_uuids, invalid_uuids = [], []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except:
                invalid_uuids.append(u)

        queryset = VisaEligibilityType.objects.filter(uuid__in=valid_uuids)
        count = queryset.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching VisaEligibilityType found.",
                "invalid_uuids": invalid_uuids
            }, status=404)

        queryset.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} VisaEligibilityType(s) deleted.",
            "invalid_uuids": invalid_uuids
        })



class VisaEligibilityTypeExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx')
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Visa Eligibility Type',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = VisaEligibilityType.objects.filter(is_deleted=False)

        if uuids:
            queryset = queryset.filter(uuid__in=uuids)

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]

        for item in queryset:
            row = []
            for field in field_list:
                value = getattr(item, field)
                row.append(value)
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'visa_eligibility_type.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'visa_eligibility_type.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


class VisaEligibilityTypeImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({"statusCode": 400, "status": False,"error": "No file uploaded"}, status=400)

        format_type = file.name.split(".")[-1].lower()
        required_headers = {"visa eligibility type"}
        optional_headers = {"description"}
        all_headers = required_headers.union(optional_headers)

        data = []
        duplicates = []
        skipped_rows = []
        to_create = []

        try:
            # ---------------- XLSX Handling ----------------
            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({"statusCode": 400, "status": False, "error": "Sheet name required", "available_sheets": available_sheets}, status=400)
                if sheet_name not in available_sheets:
                    return Response({"statusCode": 400, "status": False,"error": f'Sheet "{sheet_name}" not found', "available_sheets": available_sheets}, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({"statusCode": 400, "status": False, "message": f'Sheet "{sheet_name}" is empty'}, status=400)

                headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
                if not required_headers.issubset(set(headers)):
                    return Response({"statusCode": 400, "status": False, "message": f"Missing required headers: {required_headers}. Found: {set(headers)}"}, status=400)

                for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = row_no
                    data.append(row_dict)

            # ---------------- CSV Handling ----------------
            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                reader = csv.DictReader(io.StringIO(decoded_file))
                for row_no, row in enumerate(reader, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = row_no
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({"statusCode": 400, "status": False, "message": f"Missing required headers: {required_headers}. Found: {set(row_lower.keys())}"}, status=400)
                    data.append(row_lower)

            else:
                return Response({"statusCode": 400, "status": False, "message": "Unsupported file type. Use .xlsx or .csv"}, status=400)

            imported_count = 0
            seen_in_file = set()

            # Fetch existing records to reduce DB hits
            existing_map = {v.name.lower(): v for v in VisaEligibilityType.objects.all()}

            # ---------------- Process Rows ----------------
            for row in reversed(data):
                row_no = row.get("_row_number", "Unknown")
                name = str(row.get("visa eligibility type") or "").strip()
                description = str(row.get("description") or "").strip()

                key_in_file = name.lower()

                # Missing required field
                if not name:
                    skipped_rows.append({
                        "Row": row_no,
                        "Visa Eligibility Type": name,
                        "Description": description,
                        "Reason": "Missing required field: visa eligibility type"
                    })
                    continue

                # Duplicate in uploaded file
                if key_in_file in seen_in_file:
                    duplicates.append({
                        "Row": row_no,
                        "Visa Eligibility Type": name,
                        "Description": description,
                        "Reason": "Duplicate in uploaded file"
                    })
                    continue
                seen_in_file.add(key_in_file)

                # Check existing in DB
                existing = existing_map.get(key_in_file)
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "Row": row_no,
                            "Visa Eligibility Type": name,
                            "Description": description,
                            "Reason": "Already exists in database"
                        })
                        continue
                    # Reactivate deleted record
                    existing.description = description
                    existing.is_deleted = False
                    existing.save()
                    imported_count += 1
                    continue

                # Prepare for bulk create
                to_create.append(VisaEligibilityType(
                    name=name,
                    description=description,
                    is_deleted=False
                ))

            # ---------------- Bulk Create ----------------
            if to_create:
                batch_size = 500
                for i in range(0, len(to_create), batch_size):
                    VisaEligibilityType.objects.bulk_create(to_create[i:i + batch_size])
                imported_count += len(to_create)

        except Exception as e:
            return Response({"statusCode": 400, "status": False, "message": str(e)}, status=400)

        # ---------------- Response ----------------
        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": list(reversed(duplicates)),
            "skipped_rows": list(reversed(skipped_rows)),
        }, status=200)



class VisaStatusListAPIView(APIView):

    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = VisaStatus.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)

        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = VisaStatusSerializer(result_page, many=True)

        return paginator.get_paginated_response(serializer.data)



class VisaStatusCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = VisaStatus.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "VisaStatus with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = VisaStatusSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "VisaStatus created successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages)
            }, status=status.HTTP_400_BAD_REQUEST)

class VisaStatusUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = VisaStatus.objects.get(uuid=uuid, is_deleted=False)
        except VisaStatus.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "VisaStatus not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = VisaStatusSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "VisaStatus updated successfully",
                "data": serializer.data
            })
        else:
            messages = []
            for field, msgs in serializer.errors.items():
                messages.extend(msgs)
            return Response({
                "statusCode": 400,
                "status": False,
                "message": " ".join(messages),
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

class VisaStatusRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = VisaStatus.objects.get(uuid=uuid, is_deleted=False)
        except VisaStatus.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "VisaStatus not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = VisaStatusSerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "VisaStatus retrieved successfully",
            "data": serializer.data
        })


class VisaStatusDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = VisaStatus.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "VisaStatus permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except VisaStatus.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "VisaStatus not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            objs = VisaStatus.objects.all()
            count = objs.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No VisaStatus found to delete.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} VisaStatus permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids = []
        invalid_uuids = []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        if not valid_uuids:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No valid UUIDs provided.",
                "data": {"invalid_uuids": invalid_uuids}
            }, status=status.HTTP_400_BAD_REQUEST)

        objs = VisaStatus.objects.filter(uuid__in=valid_uuids)
        count = objs.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching VisaStatus found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=status.HTTP_404_NOT_FOUND)

        objs.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} VisaStatus permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })



class VisaStatusExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Visa Status',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        if fields:
            field_list = [f.strip() for f in fields.split(',')]
        else:
            field_list = list(field_header_map.keys())

        queryset = VisaStatus.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]
        dataset.title = 'VisaStatus'

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'visastatus.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'visastatus.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


class VisaStatusImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No file uploaded"
            }, status=400)

        format_type = file.name.split(".")[-1].lower()
        required_headers = {"visa status"}
        optional_headers = {"description"}

        data = []
        duplicates = []
        skipped_rows = []
        to_create = []

        try:
            # ---------------- XLSX Handling ----------------
            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": "Please provide sheet_name",
                        "available_sheets": available_sheets
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets
                    }, status=400)

                ws = wb[sheet_name]
                if ws.max_row <= 1:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" is empty.'
                    }, status=400)

                headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            # ---------------- CSV Handling ----------------
            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                reader = csv.DictReader(io.StringIO(decoded_file))
                for idx, row in enumerate(reader, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}"
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv"
                }, status=400)

            imported_count = 0
            seen_in_file = set()
            existing_map = {v.name.lower(): v for v in VisaStatus.objects.all()}

            # ---------------- Process Rows ----------------
            for row in reversed(data):
                row_no = row.get("_row_number", "Unknown")
                name = str(row.get("visa status") or "").strip()
                description = str(row.get("description") or "").strip()
                key_in_file = name.lower()

                # Missing required field
                if not name:
                    skipped_rows.append({
                        "Row": row_no,
                        "Visa Status": name,
                        "Description": description,
                        "Reason": "Missing required field: visa status"
                    })
                    continue

                # Duplicate in uploaded file
                if key_in_file in seen_in_file:
                    duplicates.append({
                        "Row": row_no,
                        "Visa Status": name,
                        "Description": description,
                        "Reason": "Duplicate in uploaded file"
                    })
                    continue
                seen_in_file.add(key_in_file)

                # Check existing in DB
                existing = existing_map.get(key_in_file)
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "Row": row_no,
                            "Visa Status": name,
                            "Description": description,
                            "Reason": "Already exists in database"
                        })
                        continue
                    # Reactivate deleted record
                    existing.description = description
                    existing.is_deleted = False
                    existing.save()
                    imported_count += 1
                    continue

                # Prepare for bulk create
                to_create.append(VisaStatus(
                    name=name,
                    description=description,
                    is_deleted=False
                ))

            # ---------------- Bulk Create ----------------
            if to_create:
                batch_size = 500
                for i in range(0, len(to_create), batch_size):
                    VisaStatus.objects.bulk_create(to_create[i:i + batch_size])
                imported_count += len(to_create)

        except Exception as e:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": str(e)
            }, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": list(reversed(duplicates)),
            "skipped_rows": list(reversed(skipped_rows))
        }, status=200)





class PossibilityLevelListAPIView(APIView):

    def get(self, request):
        search = request.GET.get('search', '').strip()
        sort_by = request.GET.get('sortBy', 'created_at')
        sort_order = request.GET.get('sortOrder', 'desc')

        allowed_sort_fields = ['name', 'description', 'updated_at']
        if sort_by not in allowed_sort_fields:
            sort_by = 'created_at'

        if sort_order == 'desc':
            sort_by = f'-{sort_by}'

        queryset = PossibilityLevel.objects.filter(is_deleted=False)

        if search:
            queryset = queryset.filter(Q(name__istartswith=search))

        queryset = queryset.order_by(sort_by)

        paginator = CustomPagination()
        result_page = paginator.paginate_queryset(queryset, request)
        serializer = PossibilityLevelSerializer(result_page, many=True)

        return paginator.get_paginated_response(serializer.data)



class PossibilityLevelCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        name = request.data.get("name", "").strip()
        existing = PossibilityLevel.objects.filter(name__iexact=name, is_deleted=False).first()

        if existing:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "PossibilityLevel with this name already exists."
            }, status=status.HTTP_400_BAD_REQUEST)

        serializer = PossibilityLevelSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "PossibilityLevel created successfully",
                "data": serializer.data
            })
        else:
            errors = " ".join([msg for msgs in serializer.errors.values() for msg in msgs])
            return Response({
                "statusCode": 400,
                "status": False,
                "message": errors
            }, status=status.HTTP_400_BAD_REQUEST)



class PossibilityLevelUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def put(self, request, uuid):
        try:
            obj = PossibilityLevel.objects.get(uuid=uuid, is_deleted=False)
        except PossibilityLevel.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "PossibilityLevel not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = PossibilityLevelSerializer(obj, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": "PossibilityLevel updated successfully",
                "data": serializer.data
            })
        else:
            errors = " ".join([msg for msgs in serializer.errors.values() for msg in msgs])
            return Response({
                "statusCode": 400,
                "status": False,
                "message": errors,
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)


class PossibilityLevelRetrieveAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request, uuid):
        try:
            obj = PossibilityLevel.objects.get(uuid=uuid, is_deleted=False)
        except PossibilityLevel.DoesNotExist:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "PossibilityLevel not found",
                "data": None
            }, status=status.HTTP_404_NOT_FOUND)

        serializer = PossibilityLevelSerializer(obj)
        return Response({
            "statusCode": 200,
            "status": True,
            "message": "PossibilityLevel retrieved successfully",
            "data": serializer.data
        })



class PossibilityLevelDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def delete(self, request, uuid=None):
        ids = request.data.get('id', None)

        if uuid:
            try:
                obj = PossibilityLevel.objects.get(uuid=uuid)
                obj.delete()
                return Response({
                    "statusCode": 204,
                    "status": True,
                    "message": "PossibilityLevel permanently deleted.",
                    "data": None
                }, status=status.HTTP_204_NO_CONTENT)
            except PossibilityLevel.DoesNotExist:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "PossibilityLevel not found.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)

        if ids == "all":
            objs = PossibilityLevel.objects.all()
            count = objs.count()
            if count == 0:
                return Response({
                    "statusCode": 404,
                    "status": False,
                    "message": "No PossibilityLevel found to delete.",
                    "data": None
                }, status=status.HTTP_404_NOT_FOUND)
            objs.delete()
            return Response({
                "statusCode": 200,
                "status": True,
                "message": f"All {count} PossibilityLevel permanently deleted.",
                "data": None
            })

        if not ids or not isinstance(ids, list):
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "Provide a list of UUIDs in 'id' field or 'all'.",
                "data": None
            }, status=status.HTTP_400_BAD_REQUEST)

        valid_uuids, invalid_uuids = [], []
        for u in ids:
            try:
                valid_uuids.append(UUID(u))
            except ValueError:
                invalid_uuids.append(u)

        queryset = PossibilityLevel.objects.filter(uuid__in=valid_uuids)
        count = queryset.count()

        if count == 0:
            return Response({
                "statusCode": 404,
                "status": False,
                "message": "No matching PossibilityLevel found.",
                "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
            }, status=status.HTTP_404_NOT_FOUND)

        queryset.delete()

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f"{count} PossibilityLevel(s) permanently deleted.",
            "data": {"invalid_uuids": invalid_uuids} if invalid_uuids else None
        })



class PossibilityLevelExportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def get(self, request):
        format_type = request.GET.get('format', 'xlsx').lower()
        fields = request.GET.get('fields')
        uuids_param = request.GET.get('uuids', '')
        uuids = [u.strip() for u in uuids_param.split(',') if u]

        field_header_map = {
            'uuid': 'UUID',
            'name': 'Possibility Level',
            'description': 'Description',
            'is_deleted': 'Deleted',
            'updated_at': 'Modified On',
        }

        field_list = [f.strip() for f in fields.split(',')] if fields else list(field_header_map.keys())

        queryset = PossibilityLevel.objects.filter(is_deleted=False)
        if uuids:
            queryset = queryset.filter(uuid__in=uuids)
        queryset = queryset.order_by('-created_at')

        dataset = Dataset()
        dataset.headers = [field_header_map.get(f, f) for f in field_list]

        for obj in queryset:
            row = []
            for field in field_list:
                value = getattr(obj, field, '')
                if field in ['created_at', 'updated_at'] and value:
                    value = timezone.localtime(value, india_tz).strftime("%d-%m-%Y %I:%M:%S %p")
                elif isinstance(value, bool):
                    value = int(value)
                row.append(value if value is not None else '')
            dataset.append(row)

        if format_type == 'csv':
            file_data = dataset.export('csv')
            content_type = 'text/csv'
            file_name = 'possibility_level.csv'
        else:
            file_data = io.BytesIO(dataset.export('xlsx'))
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            file_name = 'possibility_level.xlsx'

        response = HttpResponse(
            file_data if format_type == 'csv' else file_data.getvalue(),
            content_type=content_type
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response


class PossibilityLevelImportAPIView(APIView):
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request):
        file = request.FILES.get("file")
        sheet_name = request.data.get("sheet_name")

        if not file:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": "No file uploaded"
            }, status=400)

        format_type = file.name.split(".")[-1].lower()
        required_headers = {"possibility level"}
        optional_headers = {"description"}

        data = []
        duplicates = []
        skipped_rows = []
        to_create = []

        try:
            # ---------------- XLSX Handling ----------------
            if format_type == "xlsx":
                wb = openpyxl.load_workbook(file, read_only=True)
                available_sheets = wb.sheetnames

                if not sheet_name:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": "Provide sheet_name",
                        "available_sheets": available_sheets
                    }, status=400)

                if sheet_name not in available_sheets:
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f'Sheet "{sheet_name}" not found',
                        "available_sheets": available_sheets
                    }, status=400)

                ws = wb[sheet_name]
                headers = [str(c.value).lower().strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

                if not required_headers.issubset(set(headers)):
                    return Response({
                        "statusCode": 400,
                        "status": False,
                        "message": f"Missing required headers. Required: {required_headers}, Found: {set(headers)}"
                    }, status=400)

                for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue
                    row_dict = dict(zip(headers, row))
                    row_dict["_row_number"] = idx
                    data.append(row_dict)

            # ---------------- CSV Handling ----------------
            elif format_type == "csv":
                decoded_file = file.read().decode("utf-8")
                reader = csv.DictReader(io.StringIO(decoded_file))
                for idx, row in enumerate(reader, start=2):
                    row_lower = {k.strip().lower(): v for k, v in row.items()}
                    row_lower["_row_number"] = idx
                    if not required_headers.issubset(set(row_lower.keys())):
                        return Response({
                            "statusCode": 400,
                            "status": False,
                            "message": f"Missing required headers. Required: {', '.join(required_headers)}"
                        }, status=400)
                    data.append(row_lower)

            else:
                return Response({
                    "statusCode": 400,
                    "status": False,
                    "message": "Unsupported file format. Use .xlsx or .csv"
                }, status=400)

            imported_count = 0
            seen_in_file = set()
            existing_map = {v.name.lower(): v for v in PossibilityLevel.objects.all()}

            # ---------------- Process Rows ----------------
            for row in reversed(data):
                row_no = row.get("_row_number", "Unknown")
                name = str(row.get("possibility level") or "").strip()
                description = str(row.get("description") or "").strip()
                key_in_file = name.lower()

                # Missing required field
                if not name:
                    skipped_rows.append({
                        "Row": row_no,
                        "Possibility Level": name,
                        "Description": description,
                        "Reason": "Missing required field: possibility level"
                    })
                    continue

                # Duplicate in uploaded file
                if key_in_file in seen_in_file:
                    duplicates.append({
                        "Row": row_no,
                        "Possibility Level": name,
                        "Description": description,
                        "Reason": "Duplicate in uploaded file"
                    })
                    continue
                seen_in_file.add(key_in_file)

                # Check existing in DB
                existing = existing_map.get(key_in_file)
                if existing:
                    if not existing.is_deleted:
                        duplicates.append({
                            "Row": row_no,
                            "Possibility Level": name,
                            "Description": description,
                            "Reason": "Already exists in database"
                        })
                        continue
                    # Reactivate deleted record
                    existing.description = description
                    existing.is_deleted = False
                    existing.save()
                    imported_count += 1
                    continue

                # Prepare for bulk create
                to_create.append(PossibilityLevel(
                    name=name,
                    description=description,
                    is_deleted=False
                ))

            # ---------------- Bulk Create ----------------
            if to_create:
                batch_size = 500
                for i in range(0, len(to_create), batch_size):
                    PossibilityLevel.objects.bulk_create(to_create[i:i + batch_size])
                imported_count += len(to_create)

        except Exception as e:
            return Response({
                "statusCode": 400,
                "status": False,
                "message": str(e)
            }, status=400)

        return Response({
            "statusCode": 200,
            "status": True,
            "message": f'Sheet "{sheet_name}" imported successfully' if sheet_name else "Import successful",
            "imported_count": imported_count,
            "duplicates": list(reversed(duplicates)),
            "skipped_rows": list(reversed(skipped_rows))
        }, status=200)


